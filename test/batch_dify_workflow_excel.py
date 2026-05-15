#!/usr/bin/env python3
"""批量调用 Dify Workflow 处理 Excel 测试用例。"""

from __future__ import annotations

import argparse
import getpass
import json
import os
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook

DEFAULT_SHEET = "汇总40"
RESULT_COLUMNS = {
    "actual_level": "实际判定的隐患级别",
    "level_correct": "级别是否判定正确",
    "criteria": "命中的重大事故隐患判定标准",
    "reason": "判定依据",
    "confidence": "置信度",
}
OUTPUT_FIELD_ALIASES = {
    "current_level": ["current_level", "当前隐患级别"],
    "actual_level": ["actual_level", "实际判定的隐患级别"],
    "level_correct": ["level_correct", "级别是否判定正确"],
    "criteria": ["criteria", "matched_criteria", "命中的重大事故隐患判定标准"],
    "reason": ["reason", "判定依据"],
    "confidence": ["confidence", "置信度"],
}
SUPPORTED_EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm"}


@dataclass
class RuntimeConfig:
    excel_path: Path | None = None
    sheet: str = DEFAULT_SHEET
    output: Path | None = None
    base_url: str = ""
    api_key: str = ""
    user: str = "batch-test-local"
    timeout: int = 60
    limit: int | None = None
    concurrency: int = 1
    description_input_key: str = "content"
    level_input_key: str = "level"
    enterprise_input_key: str = "category"
    interactive: bool = False


@dataclass
class BatchContext:
    workbook: Any
    worksheet: Any
    source_columns: dict[str, int]
    rows: list[tuple[int, dict[str, str]]]
    output_path: Path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="读取 Excel 测试用例，逐行调用 Dify Workflow，并将结果写入新 Excel。"
    )
    parser.add_argument("excel_path", nargs="?", help="输入 Excel 文件路径")
    parser.add_argument("--sheet", default=DEFAULT_SHEET, help=f"要处理的 sheet，默认 {DEFAULT_SHEET}")
    parser.add_argument("--output", help="输出 Excel 文件路径；默认在原文件旁生成 *_结果.xlsx")
    parser.add_argument(
        "--base-url",
        default=os.getenv("DIFY_BASE_URL"),
        help="Dify Base URL，可用环境变量 DIFY_BASE_URL",
    )
    parser.add_argument(
        "--api-key",
        default=os.getenv("DIFY_API_KEY"),
        help="Dify API Key，可用环境变量 DIFY_API_KEY",
    )
    parser.add_argument(
        "--user",
        default=os.getenv("DIFY_USER", "batch-test-local"),
        help="Dify user 字段，默认 batch-test-local",
    )
    parser.add_argument("--timeout", type=int, default=60, help="单次请求超时时间（秒），默认 60")
    parser.add_argument("--limit", type=int, help="只处理前 N 条非空记录，便于调试")
    parser.add_argument("--concurrency", type=int, default=1, help="并发请求数，默认 1")
    parser.add_argument(
        "--description-input-key",
        default=os.getenv("DIFY_DESCRIPTION_INPUT_KEY", "content"),
        help="workflow 中“隐患描述”对应的输入字段名，默认 content",
    )
    parser.add_argument(
        "--level-input-key",
        default=os.getenv("DIFY_LEVEL_INPUT_KEY", "level"),
        help="workflow 中“隐患级别”对应的输入字段名，默认 level",
    )
    parser.add_argument(
        "--enterprise-input-key",
        default=os.getenv("DIFY_ENTERPRISE_INPUT_KEY", "category"),
        help="workflow 中“企业类型”对应的输入字段名，输入空字符串可跳过该字段，默认 category",
    )
    parser.add_argument("--interactive", action="store_true", help="强制进入交互式配置向导")
    return parser.parse_args()


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def clean_input_text(value: str) -> str:
    return value.strip().strip('"').strip("'")


def to_optional_path(value: str | None) -> Path | None:
    cleaned = clean_input_text(value or "")
    if not cleaned:
        return None
    return Path(cleaned).expanduser().resolve()


def mask_api_key(api_key: str) -> str:
    if len(api_key) <= 8:
        return "*" * len(api_key)
    return f"{api_key[:4]}***{api_key[-4:]}"


def build_output_path(input_path: Path, output_path: str | None) -> Path:
    if output_path:
        return Path(output_path).expanduser().resolve()

    candidate = input_path.with_name(f"{input_path.stem}_结果{input_path.suffix}")
    if not candidate.exists():
        return candidate

    timestamp = time.strftime("%Y%m%d_%H%M%S")
    return input_path.with_name(f"{input_path.stem}_结果_{timestamp}{input_path.suffix}")


def load_workbook_and_sheet(excel_path: Path, sheet_name: str):
    workbook = load_workbook(excel_path)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"sheet 不存在：{sheet_name}，可选值：{', '.join(workbook.sheetnames)}")
    return workbook, workbook[sheet_name]


def list_sheet_names(excel_path: Path) -> list[str]:
    workbook = load_workbook(excel_path, read_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        close_method = getattr(workbook, "close", None)
        if callable(close_method):
            close_method()


def resolve_source_columns(worksheet) -> dict[str, int]:
    headers = [normalize_text(cell.value) for cell in worksheet[1]]
    header_map = {name: index for index, name in enumerate(headers, start=1) if name}

    missing = [name for name in ["隐患描述", "隐患级别"] if name not in header_map]
    if missing:
        raise ValueError(f"缺少必需列：{', '.join(missing)}")

    return {
        "hazard_description": header_map["隐患描述"],
        "hazard_level": header_map["隐患级别"],
        "enterprise_type": header_map.get("企业类型", 0),
    }


def prepare_output_columns(worksheet) -> dict[str, int]:
    headers = [normalize_text(cell.value) for cell in worksheet[1]]
    header_map = {name: index for index, name in enumerate(headers, start=1) if name}
    next_column = len(headers) + 1
    output_map: dict[str, int] = {}

    for field_name, column_name in RESULT_COLUMNS.items():
        if column_name in header_map:
            output_map[field_name] = header_map[column_name]
            continue
        worksheet.cell(row=1, column=next_column, value=column_name)
        output_map[field_name] = next_column
        next_column += 1

    return output_map


def iter_rows(worksheet, source_columns: dict[str, int], limit: int | None):
    processed = 0
    for row_index in range(2, worksheet.max_row + 1):
        hazard_description = normalize_text(
            worksheet.cell(row=row_index, column=source_columns["hazard_description"]).value
        )
        hazard_level = normalize_text(
            worksheet.cell(row=row_index, column=source_columns["hazard_level"]).value
        )
        enterprise_type = ""
        if source_columns["enterprise_type"]:
            enterprise_type = normalize_text(
                worksheet.cell(row=row_index, column=source_columns["enterprise_type"]).value
            )

        if not any([hazard_description, hazard_level, enterprise_type]):
            continue

        yield row_index, {
            "hazard_description": hazard_description,
            "hazard_level": hazard_level,
            "enterprise_type": enterprise_type,
        }
        processed += 1
        if limit is not None and processed >= limit:
            break


def maybe_parse_json(value: Any) -> Any:
    if not isinstance(value, str):
        return value

    stripped = value.strip()
    if not stripped or stripped[0] not in "[{":
        return value

    try:
        return json.loads(stripped)
    except json.JSONDecodeError:
        return value


def has_output_keys(container: dict[str, Any]) -> bool:
    keys = set(container.keys())
    aliases = {alias for names in OUTPUT_FIELD_ALIASES.values() for alias in names}
    return bool(keys & aliases)


def find_structured_output(value: Any, depth: int = 0) -> dict[str, Any] | None:
    if depth > 4:
        return None

    parsed = maybe_parse_json(value)
    if isinstance(parsed, dict):
        if has_output_keys(parsed):
            return parsed
        for nested in parsed.values():
            found = find_structured_output(nested, depth + 1)
            if found is not None:
                return found
        return parsed if depth == 0 else None

    if isinstance(parsed, list):
        for item in parsed:
            found = find_structured_output(item, depth + 1)
            if found is not None:
                return found

    return None


def pick_first_value(container: dict[str, Any], aliases: list[str]) -> Any:
    for alias in aliases:
        if alias in container and container[alias] not in (None, ""):
            return container[alias]
    return None


def normalize_bool(value: Any) -> bool | None:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        if value == 1:
            return True
        if value == 0:
            return False
    if isinstance(value, str):
        normalized = value.strip().lower()
        if normalized in {"true", "1", "yes", "y", "是", "正确"}:
            return True
        if normalized in {"false", "0", "no", "n", "否", "错误"}:
            return False
    return None


def normalize_confidence(value: Any) -> Any:
    if value in (None, ""):
        return ""
    try:
        return float(value)
    except (TypeError, ValueError):
        return normalize_text(value)


def format_matched_criteria(value: Any) -> str:
    parsed = maybe_parse_json(value)

    if parsed in (None, ""):
        return ""
    if isinstance(parsed, str):
        return parsed.strip()
    if isinstance(parsed, list):
        items = [format_matched_criteria(item) for item in parsed]
        return "\n".join(item for item in items if item)
    if isinstance(parsed, dict):
        preferred_keys = [
            "category",
            "standard_source",
            "criterion",
            "standard_explanation",
        ]
        parts = [normalize_text(parsed.get(key)) for key in preferred_keys if normalize_text(parsed.get(key))]
        if parts:
            return " | ".join(parts)
        return json.dumps(parsed, ensure_ascii=False)
    return normalize_text(parsed)


def parse_workflow_result(response_json: dict[str, Any], source_level: str) -> dict[str, Any]:
    data = response_json.get("data") if isinstance(response_json, dict) else None
    status = normalize_text(data.get("status") if isinstance(data, dict) else "")
    error_message = normalize_text(data.get("error") if isinstance(data, dict) else "")

    candidates = []
    if isinstance(data, dict):
        candidates.extend([
            data.get("outputs"),
            data.get("output"),
            data.get("answer"),
            data,
        ])
    if isinstance(response_json, dict):
        candidates.extend([
            response_json.get("outputs"),
            response_json.get("output"),
            response_json.get("answer"),
            response_json,
        ])

    output_container: dict[str, Any] = {}
    for candidate in candidates:
        structured = find_structured_output(candidate)
        if structured is not None:
            output_container = structured
            if has_output_keys(structured):
                break

    current_level = normalize_text(
        pick_first_value(output_container, OUTPUT_FIELD_ALIASES["current_level"])
    ) or source_level
    actual_level = normalize_text(
        pick_first_value(output_container, OUTPUT_FIELD_ALIASES["actual_level"])
    )
    level_correct = normalize_bool(
        pick_first_value(output_container, OUTPUT_FIELD_ALIASES["level_correct"])
    )
    if level_correct is None and current_level and actual_level:
        level_correct = current_level == actual_level

    criteria = format_matched_criteria(
        pick_first_value(output_container, OUTPUT_FIELD_ALIASES["criteria"])
    )
    reason = normalize_text(
        pick_first_value(output_container, OUTPUT_FIELD_ALIASES["reason"])
    )
    confidence = normalize_confidence(
        pick_first_value(output_container, OUTPUT_FIELD_ALIASES["confidence"])
    )

    if not reason and error_message:
        reason = error_message
    if not reason and status and status.lower() != "succeeded":
        reason = f"workflow 状态异常：{status}"

    return {
        "actual_level": actual_level,
        "level_correct": level_correct if level_correct is not None else "",
        "criteria": criteria,
        "reason": reason,
        "confidence": confidence,
    }


def build_request_inputs(row_data: dict[str, str], config: RuntimeConfig) -> dict[str, str]:
    inputs = {
        config.description_input_key: row_data["hazard_description"],
        config.level_input_key: row_data["hazard_level"],
    }
    if config.enterprise_input_key and row_data.get("enterprise_type"):
        inputs[config.enterprise_input_key] = row_data["enterprise_type"]
    return inputs


def call_dify_workflow(
    base_url: str,
    api_key: str,
    inputs: dict[str, str],
    user: str,
    timeout: int,
) -> dict[str, Any]:
    response = requests.post(
        f"{base_url.rstrip('/')}/workflows/run",
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        json={
            "inputs": inputs,
            "response_mode": "blocking",
            "user": user,
        },
        timeout=timeout,
    )
    response.raise_for_status()
    try:
        return response.json()
    except ValueError as exc:
        raise RuntimeError(f"响应不是合法 JSON：{response.text[:200]}") from exc


def make_failure_result(message: str) -> dict[str, Any]:
    return {
        "actual_level": "",
        "level_correct": "",
        "criteria": "",
        "reason": message,
        "confidence": "",
    }


def format_request_exception(exc: requests.RequestException) -> str:
    message = f"HTTP 请求失败：{exc}"
    response = getattr(exc, "response", None)
    if response is not None:
        body = normalize_text(response.text)
        if body:
            if len(body) > 500:
                body = f"{body[:500]}..."
            message = f"{message} | 响应: {body}"
    return message


def process_row(row_index: int, row_data: dict[str, str], config: RuntimeConfig) -> tuple[int, dict[str, Any], bool]:
    if not row_data["hazard_description"]:
        return row_index, make_failure_result("调用失败：隐患描述为空"), False

    try:
        inputs = build_request_inputs(row_data, config)
        response_json = call_dify_workflow(
            base_url=config.base_url,
            api_key=config.api_key,
            inputs=inputs,
            user=config.user,
            timeout=config.timeout,
        )
        return row_index, parse_workflow_result(response_json, row_data["hazard_level"]), True
    except requests.RequestException as exc:
        return row_index, make_failure_result(format_request_exception(exc)), False
    except Exception as exc:
        return row_index, make_failure_result(f"处理失败：{exc}"), False


def write_result_row(worksheet, row_index: int, output_columns: dict[str, int], result: dict[str, Any]) -> None:
    for field_name, column_index in output_columns.items():
        worksheet.cell(row=row_index, column=column_index, value=result.get(field_name, ""))


def prompt_text(prompt: str, default: str | None = None, allow_blank: bool = False, blank_token: str | None = None) -> str:
    while True:
        suffix = f" [{default}]" if default not in (None, "") else ""
        blank_hint = f"（输入 {blank_token} 表示留空）" if blank_token else ""
        raw = input(f"{prompt}{suffix}{blank_hint}: ").strip()
        if blank_token and raw == blank_token:
            return ""
        if raw:
            return clean_input_text(raw)
        if default not in (None, ""):
            return str(default)
        if allow_blank:
            return ""
        print("输入不能为空，请重新输入。")


def prompt_secret(prompt: str, default: str | None = None) -> str:
    if default:
        print(f"当前默认 API Key: {mask_api_key(default)}")
    while True:
        prompt_text_value = f"{prompt}（回车沿用默认值）: " if default else f"{prompt}: "
        raw = getpass.getpass(prompt_text_value).strip()
        if raw:
            return raw
        if default:
            return default
        print("API Key 不能为空，请重新输入。")


def prompt_int(prompt: str, default: int | None = None, min_value: int | None = None, allow_blank: bool = False) -> int | None:
    while True:
        suffix = f" [{default}]" if default is not None else ""
        blank_hint = "（回车表示全部）" if allow_blank else ""
        raw = input(f"{prompt}{suffix}{blank_hint}: ").strip()
        if not raw:
            if default is not None:
                return default
            if allow_blank:
                return None
            print("输入不能为空，请重新输入。")
            continue
        try:
            value = int(raw)
        except ValueError:
            print("请输入整数。")
            continue
        if min_value is not None and value < min_value:
            print(f"输入值必须大于等于 {min_value}。")
            continue
        return value


def prompt_yes_no(prompt: str, default: bool = False) -> bool:
    suffix = "[Y/n]" if default else "[y/N]"
    while True:
        raw = input(f"{prompt} {suffix}: ").strip().lower()
        if not raw:
            return default
        if raw in {"y", "yes", "是", "确认"}:
            return True
        if raw in {"n", "no", "否", "取消"}:
            return False
        print("请输入 y 或 n。")


def build_initial_config(args: argparse.Namespace) -> RuntimeConfig:
    return RuntimeConfig(
        excel_path=to_optional_path(args.excel_path),
        sheet=normalize_text(args.sheet) or DEFAULT_SHEET,
        output=to_optional_path(args.output),
        base_url=normalize_text(args.base_url),
        api_key=normalize_text(args.api_key),
        user=normalize_text(args.user) or "batch-test-local",
        timeout=args.timeout,
        limit=args.limit,
        concurrency=args.concurrency,
        description_input_key=normalize_text(args.description_input_key) or "content",
        level_input_key=normalize_text(args.level_input_key) or "level",
        enterprise_input_key=normalize_text(args.enterprise_input_key),
        interactive=bool(args.interactive),
    )


def should_use_interactive(config: RuntimeConfig) -> bool:
    return config.interactive or not (config.excel_path and config.base_url and config.api_key)


def collect_interactive_config(config: RuntimeConfig) -> RuntimeConfig:
    print("=" * 60)
    print("Dify 批量测件工具")
    print("按提示填写信息，确认后再开始执行。")
    print("=" * 60)

    while True:
        excel_default = str(config.excel_path) if config.excel_path else None
        excel_text = prompt_text("请输入 Excel 文件路径", default=excel_default)
        excel_path = to_optional_path(excel_text)
        if excel_path is None or not excel_path.exists() or not excel_path.is_file():
            print("Excel 文件不存在，请重新输入。")
            continue
        if excel_path.suffix.lower() not in SUPPORTED_EXCEL_SUFFIXES:
            print("文件后缀不是常见 Excel 格式，请确认后重新输入。")
            continue
        config.excel_path = excel_path
        break

    sheet_names = list_sheet_names(config.excel_path)
    default_sheet = config.sheet or DEFAULT_SHEET
    while True:
        sheet = prompt_text(
            f"请输入 sheet 名（可选：{', '.join(sheet_names)}）",
            default=default_sheet,
        )
        if sheet in sheet_names:
            config.sheet = sheet
            break
        print(f"sheet 不存在：{sheet}")

    while True:
        output_default = str(config.output) if config.output else None
        output_text = prompt_text(
            "输出文件路径（回车使用自动生成的 *_结果.xlsx）",
            default=output_default,
            allow_blank=True,
            blank_token="-",
        )
        if not output_text:
            config.output = None
            break
        output_path = to_optional_path(output_text)
        if output_path is None:
            config.output = None
            break
        if output_path == config.excel_path:
            print("输出文件不能与输入文件相同，请重新输入。")
            continue
        if output_path.exists() and not prompt_yes_no("输出文件已存在，是否覆盖？", default=False):
            if prompt_yes_no("是否改为自动生成新文件名？", default=True):
                config.output = None
                break
            continue
        config.output = output_path
        break

    config.base_url = prompt_text("Dify Base URL", default=config.base_url or None)
    config.api_key = prompt_secret("Dify API Key", default=config.api_key or None)
    config.user = prompt_text("Dify user", default=config.user or "batch-test-local")
    config.timeout = prompt_int("单次请求超时时间（秒）", default=config.timeout or 60, min_value=1) or 60
    config.concurrency = prompt_int("并发数", default=config.concurrency or 1, min_value=1) or 1
    config.limit = prompt_int("限制处理条数", default=config.limit, min_value=1, allow_blank=True)

    print("\n请确认 workflow 输入字段映射：")
    config.description_input_key = prompt_text("隐患描述字段名", default=config.description_input_key or "content")
    config.level_input_key = prompt_text("隐患级别字段名", default=config.level_input_key or "level")
    config.enterprise_input_key = prompt_text(
        "企业类型字段名",
        default=config.enterprise_input_key or "category",
        allow_blank=True,
        blank_token="-",
    )
    return config


def validate_runtime_config(config: RuntimeConfig) -> None:
    if config.excel_path is None:
        raise ValueError("缺少 Excel 文件路径。")
    if not config.excel_path.exists() or not config.excel_path.is_file():
        raise ValueError(f"文件不存在：{config.excel_path}")
    if config.excel_path.suffix.lower() not in SUPPORTED_EXCEL_SUFFIXES:
        raise ValueError("输入文件不是受支持的 Excel 格式。")
    if not config.base_url:
        raise ValueError("缺少 Dify Base URL。")
    if not config.api_key:
        raise ValueError("缺少 Dify API Key。")
    if not config.description_input_key:
        raise ValueError("隐患描述字段名不能为空。")
    if not config.level_input_key:
        raise ValueError("隐患级别字段名不能为空。")
    if config.timeout < 1:
        raise ValueError("超时时间必须大于等于 1。")
    if config.concurrency < 1:
        raise ValueError("并发数必须大于等于 1。")
    if config.limit is not None and config.limit < 1:
        raise ValueError("限制处理条数必须大于等于 1。")
    if config.output is not None and config.output == config.excel_path:
        raise ValueError("输出文件不能与输入文件相同。")


def prepare_batch_context(config: RuntimeConfig) -> BatchContext:
    workbook, worksheet = load_workbook_and_sheet(config.excel_path, config.sheet)
    source_columns = resolve_source_columns(worksheet)
    rows = list(iter_rows(worksheet, source_columns, config.limit))
    output_path = build_output_path(
        config.excel_path,
        str(config.output) if config.output is not None else None,
    )
    return BatchContext(
        workbook=workbook,
        worksheet=worksheet,
        source_columns=source_columns,
        rows=rows,
        output_path=output_path,
    )


def print_run_summary(config: RuntimeConfig, context: BatchContext) -> None:
    print("\n" + "=" * 60)
    print("本次执行配置确认")
    print("=" * 60)
    print(f"输入文件: {config.excel_path}")
    print(f"输出文件: {context.output_path}")
    print(f"处理 sheet: {config.sheet}")
    print(f"待处理记录数: {len(context.rows)}")
    print(f"Dify Base URL: {config.base_url.rstrip('/')}")
    print(f"API Key: {mask_api_key(config.api_key)}")
    print(f"Dify user: {config.user}")
    print(f"超时时间: {config.timeout} 秒")
    print(f"并发数: {config.concurrency}")
    print(f"处理条数限制: {config.limit if config.limit is not None else '全部'}")
    print("输入字段映射:")
    print(f"- 隐患描述 -> {config.description_input_key}")
    print(f"- 隐患级别 -> {config.level_input_key}")
    print(f"- 企业类型 -> {config.enterprise_input_key or '跳过'}")
    print("=" * 60)


def run_batch(config: RuntimeConfig, context: BatchContext) -> int:
    worksheet = context.worksheet
    output_columns = prepare_output_columns(worksheet)
    total_rows = len(context.rows)
    success_rows = 0
    failed_rows = 0

    print(f"\n开始执行，输出文件将保存到: {context.output_path}")

    if config.concurrency == 1:
        for current, (row_index, row_data) in enumerate(context.rows, start=1):
            _, result, succeeded = process_row(row_index, row_data, config)
            write_result_row(worksheet, row_index, output_columns, result)
            if succeeded:
                success_rows += 1
                print(f"[{current}/{total_rows}] 第 {row_index} 行完成")
            else:
                failed_rows += 1
                print(f"[{current}/{total_rows}] 第 {row_index} 行失败：{result['reason']}", file=sys.stderr)
    else:
        with ThreadPoolExecutor(max_workers=config.concurrency) as executor:
            future_map = {
                executor.submit(process_row, row_index, row_data, config): row_index
                for row_index, row_data in context.rows
            }
            completed = 0
            for future in as_completed(future_map):
                completed += 1
                row_index = future_map[future]
                try:
                    _, result, succeeded = future.result()
                except Exception as exc:
                    result = make_failure_result(f"处理失败：{exc}")
                    succeeded = False
                write_result_row(worksheet, row_index, output_columns, result)
                if succeeded:
                    success_rows += 1
                    print(f"[{completed}/{total_rows}] 第 {row_index} 行完成")
                else:
                    failed_rows += 1
                    print(f"[{completed}/{total_rows}] 第 {row_index} 行失败：{result['reason']}", file=sys.stderr)

    try:
        context.workbook.save(context.output_path)
    except Exception as exc:
        print(f"保存输出文件失败：{exc}", file=sys.stderr)
        return 1

    print("\n处理完成：")
    print(f"- 处理记录数: {total_rows}")
    print(f"- 成功: {success_rows}")
    print(f"- 失败: {failed_rows}")
    print(f"- 结果文件: {context.output_path}")
    return 0


def main() -> int:
    args = parse_args()
    config = build_initial_config(args)
    interactive = should_use_interactive(config)

    if interactive:
        config = collect_interactive_config(config)

    try:
        validate_runtime_config(config)
        context = prepare_batch_context(config)
    except Exception as exc:
        print(f"初始化失败：{exc}", file=sys.stderr)
        return 1

    if interactive:
        print_run_summary(config, context)
        if not prompt_yes_no("确认开始执行吗？", default=False):
            print("已取消执行。")
            return 0
    else:
        print(f"输入文件: {config.excel_path}")
        print(f"输出文件: {context.output_path}")
        print(f"处理 sheet: {config.sheet}")
        print(f"并发数: {config.concurrency}")
        print(f"Dify Base URL: {config.base_url.rstrip('/')}")
        print(f"API Key: {mask_api_key(config.api_key)}")
        print(
            "输入字段映射: "
            f"隐患描述->{config.description_input_key}, "
            f"隐患级别->{config.level_input_key}, "
            f"企业类型->{config.enterprise_input_key or '跳过'}"
        )

    return run_batch(config, context)


if __name__ == "__main__":
    raise SystemExit(main())
