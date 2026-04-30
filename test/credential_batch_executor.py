#!/usr/bin/env python3
from __future__ import annotations

import argparse
import base64
import csv
import json
import os
import re
import shutil
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

import fitz
import requests
from openpyxl import load_workbook

SUPPORTED_EXTENSIONS = {".jpg", ".jpeg", ".png", ".pdf", ".heic"}
RESULT_HEADERS = [
    "原文件名",
    "原路径",
    "识别文件类型",
    "识别证照类型",
    "是否识别成功",
    "是否写入Excel",
    "是否完成重命名",
    "重命名后文件名",
    "workbook_sheet",
    "workbook_row",
    "缺失字段",
    "失败原因",
    "任务状态",
]
SAFE_FILENAME_RE = re.compile(r'[\\/:*?"<>|\x00-\x1f]')
DATE_PATTERNS = [
    re.compile(r"(\d{4})[-./年](\d{1,2})[-./月](\d{1,2})日?")
]
YEAR_MONTH_PATTERNS = [
    re.compile(r"(\d{4})[-./年](\d{1,2})(?:月)?")
]
DATE_RANGE_SEPARATORS = ["至", "到", "-", "~", "—", "–"]
SHEET_FIELDS = {
    "主要负责人及安全管理人员": ["序号", "证号", "姓名", "人员类型", "行业类别", "初领日期", "有效期起", "有效期止", "签发机关"],
    "特种（设备）作业人员": ["序号", "证号", "姓名", "作业类别", "操作项目", "初领日期", "有效期起", "有效期止", "应复审日期", "签发机关"],
    "其他资质": ["序号", "证号", "姓名", "证书类别", "资质名称", "有效期起", "有效期止", "应复审日期", "签发机关"],
}
CREDENTIAL_TO_SHEET = {
    "安全合格证": "主要负责人及安全管理人员",
    "特种作业操作证": "特种（设备）作业人员",
    "特种（设备）作业证书": "特种（设备）作业人员",
    "其他": "其他资质",
}


@dataclass
class ModelConfig:
    base_url: str
    api_key: str
    model_name: str


class CredentialModelClient:
    def __init__(self, config: ModelConfig, prompt_text: str):
        self.config = config
        self.prompt_text = prompt_text
        self.headers = {
            "Content-Type": "application/json",
            "Authorization": f"Bearer {config.api_key}",
        }

    def recognize_image_bytes(self, image_bytes: bytes, source_file: str, file_type: str) -> dict[str, Any]:
        base64_image = base64.b64encode(image_bytes).decode("utf-8")
        prompt_variants = [
            (
                (
                    f"{self.prompt_text}\n\n"
                    f'请将 JSON 中的 source_file 固定填写为："{source_file}"。\n'
                    f'请将 JSON 中的 file_type 固定填写为："{file_type}"。\n'
                    "只输出最终 JSON，不要输出分析过程。"
                ),
                4096,
            ),
            (
                (
                    "请识别这张个人资质证照页面，只输出一个合法 JSON 对象，不要输出解释、分析过程或 markdown。\n"
                    f'JSON 中的 source_file 固定填写为："{source_file}"。\n'
                    f'JSON 中的 file_type 固定填写为："{file_type}"。\n'
                    "只提取当前页清晰可见的信息；看不清填 null；同一 PDF 的不同页后续会合并；特种设备相关证件归为“其他”。\n"
                    "JSON 结构固定为："
                    '{"source_file":"原文件名","file_type":"image|pdf","credential_type":"安全合格证|特种作业操作证|其他|未识别","success":true,"fields":{"证号":null,"姓名":null,"性别":null,"人员类型":null,"行业类别":null,"作业类别":null,"操作项目":null,"证书类别":null,"资质名称":null,"初领日期":null,"有效期限":null,"应复审日期":null,"签发机关":null},"issues":[]}'
                ),
                4096,
            ),
        ]
        last_error: Exception | None = None
        for prompt, max_tokens in prompt_variants:
            payload = {
                "model": self.config.model_name,
                "messages": [
                    {
                        "role": "user",
                        "content": [
                            {"type": "text", "text": prompt},
                            {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{base64_image}"}},
                        ],
                    }
                ],
                "temperature": 0,
                "max_tokens": max_tokens,
            }
            response = requests.post(self.config.base_url, headers=self.headers, json=payload, timeout=600)
            response.raise_for_status()
            data = response.json()
            try:
                content = self._extract_content(data)
                return self._parse_json_content(content)
            except Exception as exc:
                last_error = exc
                continue
        if last_error is not None:
            raise last_error
        raise ValueError("模型识别失败")

    def _extract_content(self, response_json: dict[str, Any]) -> str:
        choices = response_json.get("choices") or []
        if not choices:
            raise ValueError("模型响应缺少 choices")
        message = choices[0].get("message") or {}
        content = message.get("content")
        if isinstance(content, str):
            return content.strip()
        if isinstance(content, list):
            texts: list[str] = []
            for item in content:
                if isinstance(item, dict) and item.get("type") == "text":
                    texts.append(item.get("text", ""))
                elif isinstance(item, str):
                    texts.append(item)
            if texts:
                return "".join(texts).strip()
        reasoning = message.get("reasoning")
        if isinstance(reasoning, str) and isinstance(content, str):
            return content.strip()
        raise ValueError(f"模型响应 content 格式不支持: {type(content).__name__}")

    def _parse_json_content(self, content: str) -> dict[str, Any]:
        cleaned = content.strip()
        if cleaned.startswith("```"):
            cleaned = re.sub(r"^```(?:json)?", "", cleaned).strip()
            cleaned = re.sub(r"```$", "", cleaned).strip()
        try:
            return json.loads(cleaned)
        except json.JSONDecodeError:
            match = re.search(r"\{[\s\S]*\}", cleaned)
            if not match:
                raise
            return json.loads(match.group(0))


@dataclass
class FileProcessResult:
    original_name: str
    original_path: str
    file_type: str
    credential_type: str = "未识别"
    recognized: bool = False
    excel_written: bool = False
    renamed: bool = False
    renamed_name: str = ""
    workbook_sheet: str = ""
    workbook_row: str = ""
    missing_fields: str = ""
    failure_reason: str = ""
    task_status: str = "失败"

    def to_row(self) -> list[str]:
        return [
            self.original_name,
            self.original_path,
            self.file_type,
            self.credential_type,
            yes_no(self.recognized),
            yes_no(self.excel_written),
            yes_no(self.renamed),
            self.renamed_name,
            self.workbook_sheet,
            self.workbook_row,
            self.missing_fields,
            self.failure_reason,
            self.task_status,
        ]


def yes_no(value: bool) -> str:
    return "是" if value else "否"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="批量识别个人资质证照并输出 Excel、结果表和 zip")
    base_dir = Path(__file__).resolve().parent
    assets_dir = base_dir / "assets"
    parser.add_argument("zip_path", help="输入 zip 路径")
    parser.add_argument("--template", default=str(assets_dir / "个人资质导入.xlsx"), help="Excel 模板路径")
    parser.add_argument("--prompt", default=str(assets_dir / "credential_extraction_prompt.md"), help="结构化提示词路径")
    parser.add_argument("--output-dir", default=str(base_dir / "output"), help="输出目录")
    parser.add_argument("--base-url", default=os.getenv("CREDENTIAL_MODEL_BASE_URL", ""), help="模型 base url")
    parser.add_argument("--api-key", default=os.getenv("CREDENTIAL_MODEL_API_KEY", ""), help="模型 api key")
    parser.add_argument("--model-name", default=os.getenv("CREDENTIAL_MODEL_NAME", ""), help="模型名称")
    return parser.parse_args()


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def load_prompt(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def decode_zip_filename(name: str) -> str:
    try:
        return name.encode("cp437").decode("utf-8")
    except UnicodeError:
        return name



def unzip_to_workdir(zip_path: Path, workdir: Path) -> list[Path]:
    input_dir = workdir / "input"
    ensure_dir(input_dir)
    candidates: list[Path] = []
    with zipfile.ZipFile(zip_path, "r") as zf:
        for info in zf.infolist():
            decoded_name = decode_zip_filename(info.filename)
            relative_path = Path(decoded_name)
            if any(part in {"", ".", ".."} for part in relative_path.parts):
                continue
            target_path = input_dir / relative_path
            if info.is_dir():
                ensure_dir(target_path)
                continue
            ensure_dir(target_path.parent)
            with zf.open(info, "r") as source, target_path.open("wb") as target:
                shutil.copyfileobj(source, target)
            if target_path.name.startswith("."):
                continue
            if "__MACOSX" in target_path.parts:
                continue
            if target_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
                continue
            candidates.append(target_path)
    return sorted(candidates)


def file_kind(path: Path) -> str:
    return "pdf" if path.suffix.lower() == ".pdf" else "image"


def render_pdf_pages(pdf_path: Path) -> list[bytes]:
    doc = fitz.open(pdf_path)
    try:
        pages: list[bytes] = []
        for page in doc:
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            pages.append(pix.tobytes("jpg"))
        return pages
    finally:
        doc.close()


def read_image_bytes(path: Path) -> bytes:
    return path.read_bytes()


def normalize_single_date(value: str) -> str | None:
    stripped = value.strip()
    if not stripped:
        return None
    for pattern in DATE_PATTERNS:
        match = pattern.search(stripped)
        if match:
            year, month, day = match.groups()
            return f"{year}-{int(month):02d}-{int(day):02d}"
    return None


def normalize_validity_period(value: Any) -> tuple[str | None, str | None]:
    if not isinstance(value, str):
        return None, None
    stripped = value.strip()
    if not stripped:
        return None, None
    dates = []
    for pattern in DATE_PATTERNS:
        for match in pattern.finditer(stripped):
            year, month, day = match.groups()
            dates.append(f"{year}-{int(month):02d}-{int(day):02d}")
    if len(dates) >= 2:
        return dates[0], dates[1]

    year_months = []
    for pattern in YEAR_MONTH_PATTERNS:
        for match in pattern.finditer(stripped):
            year, month = match.groups()
            year_months.append(f"{year}-{int(month):02d}")
    if len(year_months) >= 2:
        return year_months[0], year_months[1]
    return None, None


def normalize_review_date(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    stripped = value.strip()
    if not stripped:
        return None
    return normalize_single_date(stripped)


def normalize_date(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    return normalize_single_date(value)


def normalize_fields(data: dict[str, Any]) -> dict[str, Any]:
    fields = data.get("fields") or {}
    normalized: dict[str, Any] = {}
    for key in [
        "序号",
        "证号",
        "姓名",
        "性别",
        "人员类型",
        "行业类别",
        "作业类别",
        "操作项目",
        "初领日期",
        "应复审日期",
        "签发机关",
        "证书类别",
        "资质名称",
    ]:
        value = fields.get(key)
        if isinstance(value, str):
            value = value.strip() or None
        if key == "初领日期":
            value = normalize_date(value)
        elif key == "应复审日期":
            value = normalize_review_date(value)
        normalized[key] = value

    valid_from, valid_to = normalize_validity_period(fields.get("有效期限"))
    normalized["有效期起"] = valid_from
    normalized["有效期止"] = valid_to
    credential_type = data.get("credential_type")
    if credential_type == "特种（设备）作业证书":
        credential_type = "特种作业操作证"
    if credential_type not in CREDENTIAL_TO_SHEET:
        credential_type = "其他"
    if credential_type == "其他":
        normalized["证书类别"] = normalized.get("证书类别") or credential_type
        normalized["资质名称"] = normalized.get("资质名称") or fields.get("资质名称") or fields.get("证件名称") or fields.get("证书名称")
    data["credential_type"] = credential_type
    data["fields"] = normalized
    issues = data.get("issues")
    if not isinstance(issues, list):
        data["issues"] = []
    return data


def choose_best_result(results: list[dict[str, Any]]) -> dict[str, Any]:
    def score(item: dict[str, Any]) -> tuple[int, int, int]:
        fields = item.get("fields") or {}
        non_empty = sum(1 for value in fields.values() if value not in (None, "", []))
        recognized = 1 if item.get("credential_type") in CREDENTIAL_TO_SHEET else 0
        successful = 1 if item.get("success") else 0
        return recognized, successful, non_empty
    return max(results, key=score)


def merge_result_fields(items: list[dict[str, Any]]) -> tuple[dict[str, Any], list[str]]:
    merged_fields: dict[str, Any] = {}
    merged_issues: list[str] = []
    for item in items:
        fields = item.get("fields") or {}
        for key, value in fields.items():
            if key not in merged_fields or merged_fields[key] in (None, "", []):
                if value not in (None, "", []):
                    merged_fields[key] = value
        for issue in item.get("issues") or []:
            if issue and issue not in merged_issues:
                merged_issues.append(issue)
    return merged_fields, merged_issues


def merge_pdf_page_results(results: list[dict[str, Any]]) -> dict[str, Any]:
    best_result = choose_best_result(results)
    recognized_types = [
        item.get("credential_type")
        for item in results
        if item.get("credential_type") in CREDENTIAL_TO_SHEET and item.get("credential_type") != "未识别"
    ]
    if not recognized_types:
        return best_result

    type_priority = {"特种作业操作证": 3, "安全合格证": 2, "其他": 1}
    type_scores: dict[str, tuple[int, int, int]] = {}
    for credential_type in set(recognized_types):
        items = [item for item in results if item.get("credential_type") == credential_type]
        merged_fields, _ = merge_result_fields(items)
        type_scores[credential_type] = (
            sum(1 for value in merged_fields.values() if value not in (None, "", [])),
            len(items),
            type_priority.get(credential_type, 0),
        )
    primary_type = max(type_scores, key=type_scores.get)
    recognized_page_results = [item for item in results if item.get("credential_type") == primary_type]
    merged_fields, merged_issues = merge_result_fields(recognized_page_results)

    operation_item = merged_fields.get("操作项目")
    issuing_authority = str(merged_fields.get("签发机关") or "")
    if primary_type == "其他" and (operation_item or "市场监督管理局" in issuing_authority):
        primary_type = "特种作业操作证"
        recognized_page_results = [
            item for item in results if item.get("credential_type") in {"特种作业操作证", "其他"}
        ]
        merged_fields, merged_issues = merge_result_fields(recognized_page_results)

    if not merged_fields:
        return best_result

    return {
        "source_file": best_result.get("source_file"),
        "file_type": best_result.get("file_type"),
        "credential_type": primary_type,
        "success": any(item.get("success") for item in recognized_page_results),
        "fields": merged_fields,
        "issues": merged_issues,
    }



def append_to_workbook(workbook, data: dict[str, Any]) -> tuple[str, int] | tuple[str, None]:
    credential_type = data["credential_type"]
    sheet_name = CREDENTIAL_TO_SHEET.get(credential_type)
    if not sheet_name or sheet_name not in SHEET_FIELDS:
        return "", None
    sheet = workbook[sheet_name]
    row = [data["fields"].get(field) for field in SHEET_FIELDS[sheet_name]]
    sheet.append(row)
    return sheet_name, sheet.max_row


def next_sheet_sequence(workbook, sheet_name: str) -> str:
    sheet = workbook[sheet_name]
    return str(sheet.max_row)


def sanitize_filename_part(value: str) -> str:
    cleaned = SAFE_FILENAME_RE.sub("_", value.strip())
    cleaned = re.sub(r"\s+", "_", cleaned)
    cleaned = re.sub(r"_+", "_", cleaned)
    return cleaned.strip("._")


def build_renamed_path(renamed_dir: Path, fields: dict[str, Any], sheet_name: str, extension: str) -> Path | None:
    sequence = fields.get("序号")
    name = fields.get("姓名")
    if not sequence or not name:
        return None

    if sheet_name == "主要负责人及安全管理人员":
        prefix = "1"
        tail_value = fields.get("人员类型")
    elif sheet_name == "特种（设备）作业人员":
        prefix = "2"
        tail_value = fields.get("操作项目")
    elif sheet_name == "其他资质":
        prefix = "3"
        tail_value = fields.get("资质名称")
    else:
        return None

    if not tail_value:
        return None

    stem = "_".join([
        f"{prefix}-{sanitize_filename_part(str(sequence))}",
        sanitize_filename_part(str(name)),
        sanitize_filename_part(str(tail_value)),
    ])
    if not stem.replace("_", ""):
        return None
    candidate = renamed_dir / f"{stem}{extension.lower()}"
    index = 2
    while candidate.exists():
        candidate = renamed_dir / f"{stem}_{index}{extension.lower()}"
        index += 1
    return candidate


def write_summary_csv(rows: list[FileProcessResult], output_path: Path) -> None:
    with output_path.open("w", encoding="utf-8-sig", newline="") as fp:
        writer = csv.writer(fp)
        writer.writerow(RESULT_HEADERS)
        for row in rows:
            writer.writerow(row.to_row())


def package_outputs(output_root: Path, workbook_path: Path, summary_path: Path, renamed_dir: Path) -> Path:
    zip_path = output_root / "个人资质批量处理结果.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.write(workbook_path, workbook_path.name)
        for file_path in renamed_dir.rglob("*"):
            if file_path.is_file():
                zf.write(file_path, file_path.name)
    return zip_path


def process_file(path: Path, client: CredentialModelClient) -> dict[str, Any]:
    kind = file_kind(path)
    if kind == "pdf":
        page_results = []
        for page_bytes in render_pdf_pages(path):
            try:
                result = client.recognize_image_bytes(page_bytes, path.name, "pdf")
            except Exception as exc:
                result = {
                    "source_file": path.name,
                    "file_type": "pdf",
                    "credential_type": "未识别",
                    "success": False,
                    "fields": {},
                    "issues": [str(exc)],
                }
            page_results.append(result)
        if not page_results:
            raise ValueError("PDF 无可识别页面")
        merged_result = merge_pdf_page_results(page_results)
        return normalize_fields(merged_result)
    result = client.recognize_image_bytes(read_image_bytes(path), path.name, "image")
    return normalize_fields(result)


def collect_missing_fields(data: dict[str, Any]) -> list[str]:
    credential_type = data.get("credential_type")
    sheet_name = CREDENTIAL_TO_SHEET.get(credential_type)
    if not sheet_name or sheet_name not in SHEET_FIELDS:
        return []
    missing = []
    fields = data.get("fields") or {}
    for field in SHEET_FIELDS[sheet_name]:
        if fields.get(field) in (None, ""):
            missing.append(field)
    return missing


def summarize_issues(issues: list[str], missing_fields: list[str]) -> str:
    filtered = []
    has_name = "姓名" not in missing_fields
    has_cert_no = "证号" not in missing_fields
    has_operation_item = "操作项目" not in missing_fields
    has_validity = "有效期起" not in missing_fields and "有效期止" not in missing_fields

    for issue in issues:
        text = str(issue).strip()
        if not text:
            continue
        if has_name and has_cert_no and ("仅封面" in text or "无具体个人信息" in text or "姓名、证号等个人信息未在当前页显示" in text):
            continue
        if has_operation_item and ("作业项目" in text and "缺失" in text or "未显示有效期、作业项目等关键信息" in text):
            continue
        if has_validity and ("有效期限缺失" in text or "无具体日期信息" in text):
            continue
        if text not in filtered:
            filtered.append(text)
    return "；".join(filtered)


def build_result_record(path: Path, data: dict[str, Any], workbook, renamed_dir: Path) -> FileProcessResult:
    record = FileProcessResult(original_name=path.name, original_path=str(path), file_type=file_kind(path))
    credential_type = data.get("credential_type", "其他")
    record.credential_type = credential_type
    record.recognized = bool(data.get("success")) and credential_type in CREDENTIAL_TO_SHEET
    issues = data.get("issues") or []

    if not record.recognized:
        missing_fields = collect_missing_fields(data)
        record.missing_fields = "、".join(missing_fields)
        record.failure_reason = summarize_issues(issues, missing_fields) or "未提取到可写入信息"
        record.task_status = "失败"
        return record

    sheet_name = CREDENTIAL_TO_SHEET.get(credential_type, "")
    if not sheet_name or sheet_name not in SHEET_FIELDS:
        record.failure_reason = "未匹配到 Excel sheet"
        record.task_status = "失败"
        return record

    data["fields"]["序号"] = next_sheet_sequence(workbook, sheet_name)
    missing_fields = collect_missing_fields(data)
    record.missing_fields = "、".join(missing_fields)

    row_sheet_name, row_no = append_to_workbook(workbook, data)
    if row_no is None:
        record.failure_reason = "未匹配到 Excel sheet"
        record.task_status = "失败"
        return record

    record.excel_written = True
    record.workbook_sheet = row_sheet_name
    record.workbook_row = str(row_no)

    target_path = build_renamed_path(renamed_dir, data["fields"], row_sheet_name, path.suffix)
    if target_path is not None:
        shutil.copy2(path, target_path)
        record.renamed = True
        record.renamed_name = target_path.name
    else:
        issues = list(issues) + ["缺少重命名必要字段"]

    if record.excel_written and record.renamed:
        record.task_status = "完成"
    elif record.excel_written:
        record.task_status = "部分完成"
    else:
        record.task_status = "失败"

    record.failure_reason = summarize_issues(issues, missing_fields)
    if missing_fields and record.task_status == "完成":
        record.task_status = "待人工核对"
    return record


def main() -> int:
    args = parse_args()
    zip_path = Path(args.zip_path).expanduser().resolve()
    template_path = Path(args.template).expanduser().resolve()
    prompt_path = Path(args.prompt).expanduser().resolve()
    output_root = Path(args.output_dir).expanduser().resolve()

    if not zip_path.exists():
        raise FileNotFoundError(f"zip 不存在: {zip_path}")
    if not template_path.exists():
        raise FileNotFoundError(f"模板不存在: {template_path}")
    if not prompt_path.exists():
        raise FileNotFoundError(f"提示词不存在: {prompt_path}")
    if not args.base_url or not args.api_key or not args.model_name:
        raise ValueError("缺少模型参数：base_url/api_key/model_name")

    ensure_dir(output_root)
    runtime_root = output_root / ".credential_batch_runtime"
    if runtime_root.exists():
        shutil.rmtree(runtime_root)
    ensure_dir(runtime_root)
    workdir = runtime_root / "workdir"
    ensure_dir(workdir)
    renamed_dir = runtime_root / "renamed_files"
    ensure_dir(renamed_dir)

    prompt_text = load_prompt(prompt_path)
    client = CredentialModelClient(ModelConfig(args.base_url, args.api_key, args.model_name), prompt_text)
    files = unzip_to_workdir(zip_path, workdir)

    workbook = load_workbook(template_path)
    results: list[FileProcessResult] = []

    for file_path in files:
        try:
            data = process_file(file_path, client)
            result = build_result_record(file_path, data, workbook, renamed_dir)
        except Exception as exc:
            result = FileProcessResult(
                original_name=file_path.name,
                original_path=str(file_path),
                file_type=file_kind(file_path),
                failure_reason=str(exc),
                task_status="失败",
            )
        results.append(result)

    workbook_path = runtime_root / "个人资质导入_识别结果.xlsx"
    workbook.save(workbook_path)

    summary_path = output_root / "个人资质_处理结果.csv"
    write_summary_csv(results, summary_path)
    zip_output = package_outputs(output_root, workbook_path, summary_path, renamed_dir)
    shutil.rmtree(runtime_root)

    summary = {
        "total_input_files": len(files),
        "recognized_counts": {
            "安全合格证": sum(1 for item in results if item.credential_type == "安全合格证"),
            "特种作业操作证": sum(1 for item in results if item.credential_type == "特种作业操作证"),
            "其他": sum(1 for item in results if item.credential_type == "其他"),
            "未识别": sum(1 for item in results if item.credential_type == "未识别"),
        },
        "processing_result": str(summary_path),
        "final_zip": str(zip_output),
        "manual_review_files": [item.original_name for item in results if item.task_status in {"部分完成", "待人工核对", "失败"}],
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
